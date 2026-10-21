"""Human interrater agreement and repeated-measures publication statistics."""
import csv
import re

import numpy as np
from scipy import stats
from openpyxl import load_workbook
from speech_pipeline import ROOT, utterances as auto
from speech_pipeline.utterances import align, flatten_labels

B = 10000
SEED = 20260924
METRIC_LABELS = {'basic_ncontent_tokens': 'Content word tokens', 'basic_ncontent_types': 'Content word types', 'basic_nfunction_tokens': 'Function word tokens', 'basic_nfunction_types': 'Function word types', 'basic_ntokens': 'Total word tokens', 'basic_ntypes': 'Total word types', 'hdd42_aw': 'HD-D-42', 'lexical_density_tokens': 'Lexical density by tokens', 'lexical_density_types': 'Lexical density by types', 'log_ttr_aw': 'Log TTR', 'maas_ttr_aw': 'Maas', 'mattr50_aw': 'MATTR-50', 'msttr50_aw': 'MSTTR-50', 'mtld_mean_factor_length_aw': 'MTLD mean factor length', 'mtld_original_aw': 'MTLD original', 'root_ttr_aw': 'Root TTR', 'simple_ttr_aw': 'TTR', 'mlu_nutterances': 'Eligible utterances', 'mlu_nwords': 'MLU word count', 'mlu_words': 'MLUw', 'mlu_words_sd': 'Utterance length SD', 'total_utterances': 'Total utterances'}


def write_csv(path, rows):
    with path.open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=list(dict.fromkeys(k for r in rows for k in r)))
        writer.writeheader()
        writer.writerows(rows)


def repeated_components(counts, sums, within_ss):
    """Profile REML for a random patient intercept and independent recording residuals.

    Vectorized over bootstrap samples. Optimize log(1 + between/within variance)
    on a grid followed by golden-section refinement, including the zero boundary.
    """
    total = counts.sum(axis=-1)
    group_means = sums/counts
    def fit(t):
        ratio = np.expm1(t)
        weights = counts/(1+counts*ratio[...,None])
        mean = (weights*group_means).sum(axis=-1)/weights.sum(axis=-1)
        q = within_ss.sum(axis=-1)+(weights*(group_means-mean[...,None])**2).sum(axis=-1)
        within = q/(total-1)
        objective = ((total-1)*np.log(np.maximum(within,1e-300))
                     +np.log1p(counts*ratio[...,None]).sum(axis=-1)
                     +np.log(weights.sum(axis=-1)))
        return objective,mean,within,ratio*within
    zero = np.zeros_like(total)
    grid=np.linspace(0,24,49)
    objectives=np.stack([fit(zero+t)[0] for t in grid])
    best=np.argmin(objectives,axis=0)
    lo=grid[np.maximum(best-1,0)]
    hi=grid[np.minimum(best+1,len(grid)-1)]
    phi=(np.sqrt(5)-1)/2
    for _ in range(55):
        left=hi-phi*(hi-lo)
        right=lo+phi*(hi-lo)
        choose_left=fit(left)[0]<fit(right)[0]
        lo,hi=np.where(choose_left,lo,left),np.where(choose_left,right,hi)
    t=(lo+hi)/2
    # For singleton-only resamples the variance partition is unidentified.
    t=np.where((fit(zero)[0]<=fit(t)[0]+1e-10)|(total==counts.shape[-1]),zero,t)
    _,mean,within,between=fit(t)
    return mean,within,between,within+between,t


def repeated_ba(rows, log_scale=False):
    x = np.asarray([r['custom_value'] for r in rows],float)
    y = np.asarray([r['reference_value'] for r in rows],float)
    if log_scale:
        if np.any(x<=0) or np.any(y<=0):
            raise ValueError('Log-ratio analysis requires strictly positive paired scores')
        x,y = np.log(x),np.log(y)
    d,m = x-y,(x+y)/2
    ids = sorted({r['patient_id'] for r in rows})
    samples = [d[[r['patient_id']==pid for r in rows]] for pid in ids]
    counts = np.array([len(v) for v in samples],float)
    sums = np.array([v.sum() for v in samples])
    within_ss = np.array([np.sum((v-v.mean())**2) for v in samples])
    bias,within,between,var,ratio_log1p = repeated_components(counts,sums,within_ss)
    z = stats.norm.ppf(.975)
    sd = np.sqrt(var)
    result = dict(bias=float(bias),lower_loa=float(bias-z*sd),upper_loa=float(bias+z*sd),
        within_patient_variance=float(within),between_patient_variance=float(between),
        variance_ratio_log1p=float(ratio_log1p),sd_difference=float(sd),
        observed_difference_min=float(d.min()),observed_difference_max=float(d.max()))
    # Sample patients with replacement, retaining every matched recording within each.
    # Each selected copy remains a distinct cluster even when its source ID is repeated.
    idx = np.random.default_rng(SEED).integers(0,len(ids),size=(B,len(ids)))
    means,_,_,variances,_ = repeated_components(counts[idx],sums[idx],within_ss[idx])
    estimates = np.stack([means,means-z*np.sqrt(variances),means+z*np.sqrt(variances)],axis=1)
    ci = np.quantile(estimates,[.025,.975],axis=0)
    zero_variation = np.ptp(d)==0
    for j,key in enumerate(['bias','lower_loa','upper_loa']):
        result[key+'_ci_lower'] = None if zero_variation else float(ci[0,j])
        result[key+'_ci_upper'] = None if zero_variation else float(ci[1,j])
    # Descriptive checks, not pass/fail tests of agreement.
    result['normality_shapiro_p'] = None if zero_variation else float(stats.shapiro(d).pvalue)
    if np.ptp(m)>0 and not zero_variation:
        slopes=[]
        indices=[np.flatnonzero([r['patient_id']==pid for r in rows]) for pid in ids]
        for draw in idx:
            take=np.concatenate([indices[i] for i in draw])
            mx,dy=m[take],d[take]
            den=np.sum((mx-mx.mean())**2)
            if den>0: slopes.append(float(np.sum((mx-mx.mean())*(dy-dy.mean()))/den))
        result['difference_mean_slope']=float(stats.linregress(m,d).slope)
        result['difference_mean_slope_ci_lower'],result['difference_mean_slope_ci_upper']=map(float,np.quantile(slopes,[.025,.975]))
    result['estimation_note'] = ('Identical observed differences. Population confidence intervals not estimated.' if zero_variation
        else 'REML between-patient variance at zero boundary.' if ratio_log1p==0 else '')
    return result


def word_stream(utterances):
    tokens, boundaries = [], set()
    for u in utterances:
        u = re.sub(r'\*unintelligible\*', ' xxx ', u, flags=re.I)
        words = [w.lower().replace('’', "'") for w in auto.WORD.findall(u)]
        # Human agreement uses a common stream despite differing filler labels.
        # Pipeline boundary comparisons report their filler policy separately.
        words = [w for w in words if not auto.is_filler(w)]
        if not words:
            continue
        tokens.extend(words)
        boundaries.add(len(tokens))
    boundaries.discard(len(tokens))
    boundaries.discard(0)
    return tokens, boundaries


def interrater_agreement(base, secondary, comparison, secondary_rater):
    a,ab=word_stream(base); b,bb=word_stream(secondary)
    mapping,edits=align(a,b)
    gaps=[]
    for i in range(len(a)-1):
        j=mapping.get(i)
        if j is not None and mapping.get(i+1)==j+1 and not ({a[i],a[i+1]} & {'xxx','yyy','www'}):
            gaps.append(dict(primary_gap=i+1,secondary_gap=j+1,left_word=a[i],right_word=a[i+1],
                primary_boundary=int(i+1 in ab),secondary_boundary=int(j+1 in bb)))
    tp=sum(r['primary_boundary'] and r['secondary_boundary'] for r in gaps)
    tn=sum(not r['primary_boundary'] and not r['secondary_boundary'] for r in gaps)
    fn=sum(r['primary_boundary'] and not r['secondary_boundary'] for r in gaps)
    fp=sum(not r['primary_boundary'] and r['secondary_boundary'] for r in gaps)
    n=len(gaps); po=(tp+tn)/n
    pe=((tp+fn)*(tp+fp)+(tn+fp)*(tn+fn))/n**2
    result=dict(comparison=comparison,secondary_rater=secondary_rater,
        filler_policy='Excluded from both raters: '+', '.join(sorted(auto.FILLERS)),
        primary_utterances=len(base),secondary_utterances=len(secondary),
        primary_tokens=len(a),secondary_tokens=len(b),primary_internal_boundaries=len(ab),secondary_internal_boundaries=len(bb),
        evaluable_gaps=n,primary_gap_coverage=n/(len(a)-1),secondary_gap_coverage=n/(len(b)-1),
        both_boundary=tp,both_no_boundary=tn,primary_only_boundary=fn,secondary_only_boundary=fp,
        boundary_f1=2*tp/(2*tp+fp+fn),boundary_precision=tp/(tp+fp),boundary_recall=tp/(tp+fn),
        boundary_percent_agreement=100*po,cohens_kappa=(po-pe)/(1-pe),
        word_edit_distance=sum(edits.values()),word_disagreement_rate=sum(edits.values())/len(a),**edits)
    return result,gaps


def pooled_interrater(rows):
    tp=sum(r['both_boundary'] for r in rows); tn=sum(r['both_no_boundary'] for r in rows)
    fn=sum(r['primary_only_boundary'] for r in rows); fp=sum(r['secondary_only_boundary'] for r in rows)
    n=tp+tn+fn+fp; po=(tp+tn)/n
    pe=((tp+fn)*(tp+fp)+(tn+fp)*(tn+fn))/n**2
    primary_tokens=sum(r['primary_tokens'] for r in rows)
    secondary_tokens=sum(r['secondary_tokens'] for r in rows)
    edits={k:sum(r[k] for r in rows) for k in ['substitutions','deletions','insertions']}
    return dict(comparison='All_primary_human_vs_secondary_raters',secondary_rater='Jade and Yamen',
        filler_policy='Excluded from both raters: '+', '.join(sorted(auto.FILLERS)),
        primary_utterances=sum(r['primary_utterances'] for r in rows),
        secondary_utterances=sum(r['secondary_utterances'] for r in rows),
        primary_tokens=primary_tokens,secondary_tokens=secondary_tokens,
        primary_internal_boundaries=sum(r['primary_internal_boundaries'] for r in rows),
        secondary_internal_boundaries=sum(r['secondary_internal_boundaries'] for r in rows),
        evaluable_gaps=n,
        primary_gap_coverage=n/sum(max(r['primary_tokens']-1,1) for r in rows),
        secondary_gap_coverage=n/sum(max(r['secondary_tokens']-1,1) for r in rows),
        both_boundary=tp,both_no_boundary=tn,primary_only_boundary=fn,secondary_only_boundary=fp,
        boundary_f1=2*tp/(2*tp+fp+fn),boundary_precision=tp/(tp+fp),boundary_recall=tp/(tp+fn),
        boundary_percent_agreement=100*po,cohens_kappa=(po-pe)/(1-pe),
        word_edit_distance=sum(edits.values()),word_disagreement_rate=sum(edits.values())/primary_tokens,**edits)


def calculate_interrater(primary_wb):
    irr_rows=[]; gap_rows=[]
    for workbook_name,rater,participants in [
        ('jade.xlsx','Jade',['P2','P9']),
        ('yamen.xlsx','Yamen',['P4','P7'])]:
        secondary_wb=load_workbook(ROOT/'dataset/annotations'/workbook_name,data_only=True)
        for participant in participants:
            primary_files=flatten_labels(primary_wb[participant])
            secondary_files=flatten_labels(secondary_wb[participant])
            for filename in sorted(set(primary_files) & set(secondary_files)):
                comparison=f'{participant}_Nafee_vs_{rater}'
                irr,gaps=interrater_agreement(primary_files[filename],secondary_files[filename],comparison,rater)
                irr['patient_id']=participant; irr['source_file']=filename
                irr_rows.append(irr)
                gap_rows.extend(dict(patient_id=participant,source_file=filename,secondary_rater=rater,**gap) for gap in gaps)
    irr_rows.append(pooled_interrater(irr_rows))
    for row in irr_rows:
        row['comparison'] = row['comparison'].replace('primary_human', 'Nafee')
        row['primary_rater'] = 'Nafee'
        row['primary_label_source'] = 'dataset/annotations/all_nafee_utterances.xlsx'
    return irr_rows,gap_rows


def main():
    primary = load_workbook(
        ROOT / 'dataset/annotations/all_nafee_utterances.xlsx',
        data_only=True,
    )
    rows, gaps = calculate_interrater(primary)
    primary.close()
    for row in rows:
        row['primary_rater'] = 'Nafee'
        row['primary_label_source'] = 'dataset/annotations/all_nafee_utterances.xlsx'
        row['comparison'] = row['comparison'].replace('primary_human', 'Nafee')
    fields = list(dict.fromkeys(key for row in rows for key in row))
    out = ROOT / 'outputs/results/paper/human_interrater_results.csv'
    with out.open('w', newline='', encoding='utf-8-sig') as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)
    print(f'Saved {len(rows) - 1} human-rater comparisons and their pooled result.')
    audit = ROOT / 'outputs/results/human_comparison/human_interrater_gaps.csv'
    audit.parent.mkdir(parents=True, exist_ok=True)
    with audit.open('w', newline='', encoding='utf-8-sig') as handle:
        writer = csv.DictWriter(handle, fieldnames=list(gaps[0]))
        writer.writeheader()
        writer.writerows(gaps)


if __name__ == "__main__":
    main()
