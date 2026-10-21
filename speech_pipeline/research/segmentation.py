"""Registered segmentation experiments: fit/select on train/validation only.

Candidate weights and feature caches stay in a temporary work directory. The
selected production bundle and auditable summaries are the project artifacts.
Run `python -m speech_pipeline.research.segmentation --help` for stages.
"""
import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
import shutil
import time

import numpy as np
import torch

from speech_pipeline import ROOT
from speech_pipeline.boundary_model import BoundaryTransformer,probabilities as legacy_probabilities
from speech_pipeline.corpora import EVIDENCE,OUT
from speech_pipeline.pipeline import save,digest
from speech_pipeline.segmentation import choose_threshold,fit_pretrained,load_bundle,metrics,pretrained_probabilities

WORK=Path('/private/tmp/speech_segmentation_research')
EXPERIMENTS=EVIDENCE/'improvement_experiments.json'


def fetch_models():
    """Download official models at immutable revisions into temporary storage."""
    from huggingface_hub import snapshot_download
    models=[('bert','google/bert_uncased_L-4_H-256_A-4','387825ce42dbb39b87911cdf8e383ee3b25184f8',
             ['config.json','vocab.txt','pytorch_model.bin']),
            ('sat','segment-any-text/sat-3l-sm','137da054051ad9f1eac42025f758db4ac9f22535',
             ['config.json','model.safetensors']),
            ('sat_tokenizer','FacebookAI/xlm-roberta-base','e73636d4f797dec63c3081bb6ed5c7b0bb3f2089',
             ['config.json','tokenizer.json','tokenizer_config.json','sentencepiece.bpe.model'])]
    for name,repo,revision,files in models:
        folder=snapshot_download(repo_id=repo,revision=revision,allow_patterns=files,cache_dir=WORK/'pretrained')
        suffix='tokenizer' if name=='sat_tokenizer' else 'model'
        pointer=f'speech_sat_{suffix}_path.txt' if name.startswith('sat') else 'speech_bert_model_path.txt'
        Path('/private/tmp',pointer).write_text(folder)
        Path('/private/tmp',f'speech_{name}_metadata.json').write_text(json.dumps(dict(id=repo,sha=revision)))
        print(name,revision,flush=True)


def load_partitions(include_test=False):
    from speech_pipeline.cohorts import partition_lookup, validate_records
    protocol=json.loads((EVIDENCE/'improvement_protocol.json').read_text())
    path=EVIDENCE/'talkbank_training.jsonl'
    if digest(path)!=protocol['input_sha256']:raise ValueError('Registered training data changed')
    groups=protocol['partitions']
    partition_lookup(groups)
    rows=[json.loads(s) for s in path.read_text().splitlines()]
    validate_records(rows,groups)
    names=['train','validation']+(['test'] if include_test else [])
    return {s:[r for r in rows if r['participant'] in groups[s]] for s in names},protocol


def record_experiment(name,entry):
    previous=json.loads(EXPERIMENTS.read_text()) if EXPERIMENTS.exists() else {}
    previous[name]=entry;save(EXPERIMENTS,previous)


def metadata(name,protocol):
    return dict(name=name,input_sha256=protocol['input_sha256'],partitions=protocol['partitions'],
        selection_scope='Training weights/vocabulary use train only; epochs/threshold/candidate use validation only',
        test_status=protocol['test_status'],code_sha256={p:digest(ROOT/p) for p in
            ['speech_pipeline/segmentation.py','speech_pipeline/research/segmentation.py']})


def cache_probabilities(name,records,compute):
    folder=WORK/'cache';folder.mkdir(parents=True,exist_ok=True)
    # Features contain no labels. Hash the actual words and derived rule flags.
    payload=[dict(tokens=r['tokens'],baseline=r['baseline']) for r in records]
    fingerprint=hashlib.sha256(json.dumps(payload,sort_keys=True).encode()).hexdigest()
    path=folder/f'{name}_{fingerprint[:16]}.npz'
    if path.exists():
        with np.load(path,allow_pickle=False) as saved:return [saved[str(i)] for i in range(len(records))]
    values=compute(records)
    np.savez_compressed(path,**{str(i):v for i,v in enumerate(values)})
    return values


def overlap_probabilities(records):
    saved=torch.load(OUT/'hybrid_model.pt',map_location='cpu',weights_only=True)
    model=BoundaryTransformer(len(saved['vocabulary'])+2,use_rules=True);model.load_state_dict(saved['state_dict'])
    chunks=[];locations=[]
    for ri,r in enumerate(records):
        for start in range(0,len(r['tokens']),128):
            end=min(start+256,len(r['tokens']))
            chunks.append(dict(tokens=r['tokens'][start:end],baseline=r['baseline'][start:end],labels=[-100]*(end-start)))
            locations.append((ri,start,end))
            if end==len(r['tokens']):break
    values=legacy_probabilities(model,chunks,saved['vocabulary'])
    totals=[np.zeros(len(r['tokens'])) for r in records];weights=[v.copy() for v in totals]
    for p,(ri,start,end) in zip(values,locations):
        n=end-start;hat=np.minimum(np.arange(n)+1,n-np.arange(n)).clip(1,64)
        logit=np.log(np.clip(p,1e-7,1-1e-7)/np.clip(1-p,1e-7,1-1e-7))
        totals[ri][start:end]+=logit*hat;weights[ri][start:end]+=hat
    result=[]
    for total,weight in zip(totals,weights):
        p=1/(1+np.exp(-np.clip(total/weight,-50,50)));p[-1]=0;result.append(p)
    return result


def sat_probabilities(records):
    from wtpsplit import SaT
    path=Path('/private/tmp/speech_sat_model_path.txt').read_text().strip()
    tokenizer=Path('/private/tmp/speech_sat_tokenizer_path.txt').read_text().strip()
    model=SaT(path,tokenizer_name_or_path=tokenizer,from_pretrained_kwargs={'local_files_only':True})
    result=[]
    for at in range(0,len(records),16):
        group=records[at:at+16];texts=[' '.join(r['tokens']) for r in group]
        values=list(model.predict_proba(texts,block_size=256,stride=128,batch_size=16,weighting='hat',outer_batch_size=16))
        for r,text,chars in zip(group,texts,values):
            ends=np.cumsum([len(w)+1 for w in r['tokens']])-2
            if len(chars)!=len(text):raise ValueError('SaT character alignment mismatch')
            p=np.asarray(chars)[ends].astype(float);p[-1]=0;result.append(p)
        if at%64==0:print(f'SaT paragraphs {min(at+16,len(records))}/{len(records)}',flush=True)
    return result


def feature_rows(records,sat,valid_only=False):
    for r,prior in zip(records,sat):
        words=r['tokens'];rules=r['baseline']
        for i in range(len(words)-1):
            if valid_only and r['labels'][i]<0:continue
            p=float(np.clip(prior[i],1e-6,1-1e-6))
            row=dict(rule=float(rules[i]),sat_logit=float(np.log(p/(1-p)))/8,
                repetition=float(words[i]==words[i+1]),left_length=min(len(words[i]),20)/20,
                right_length=min(len(words[i+1]),20)/20,
                left_suffix=words[i][-3:],right_prefix=words[i+1][:3],
                pair=words[i]+'|'+words[i+1])
            for offset in range(-2,4):
                j=i+offset;row[f'word{offset}']=words[j] if 0<=j<len(words) else '_EDGE_'
                row[f'rule{offset}']=float(rules[j]) if 0<=j<len(words)-1 else 0.
            yield row


def lexical_probabilities(records,sat,vectorizer,classifier):
    x=small_csr(vectorizer.transform(feature_rows(records,sat)))
    values=classifier.predict_proba(x)[:,1];result=[];at=0
    for r in records:
        n=len(r['tokens'])-1;result.append(np.r_[values[at:at+n],0.]);at+=n
    if at!=len(values):raise ValueError('Lexical word-gap alignment mismatch')
    return result


def small_csr(x):
    # sklearn's SGD solver requires 32-bit CSR indices even when scipy chooses
    # int64 for a small matrix. Refuse overflow rather than silently truncating.
    if x.nnz>=2**31 or x.shape[1]>=2**31:raise ValueError('Sparse feature matrix exceeds SGD index capacity')
    x.indices=x.indices.astype(np.int32,copy=False);x.indptr=x.indptr.astype(np.int32,copy=False)
    return x


def baselines():
    from sklearn.feature_extraction import DictVectorizer
    from sklearn.linear_model import SGDClassifier
    import joblib
    torch.set_num_threads(4);WORK.mkdir(parents=True,exist_ok=True)
    splits,protocol=load_partitions();train=splits['train'];validation=splits['validation']
    p=cache_probabilities('hybrid_overlap',validation,overlap_probabilities)
    choice=choose_threshold(validation,p)
    record_experiment('hybrid_overlap',dict(**metadata('hybrid_overlap',protocol),status='complete',
        selected=choice,validation=metrics(validation,p,choice['threshold']),cost='Existing 332,930-parameter checkpoint; two context views'))
    sat_train=cache_probabilities('sat',train,sat_probabilities)
    sat_validation=cache_probabilities('sat',validation,sat_probabilities)
    choice=choose_threshold(validation,sat_validation)
    record_experiment('sat_pretrained',dict(**metadata('sat_pretrained',protocol),status='complete',selected=choice,
        validation=metrics(validation,sat_validation,choice['threshold']),
        pretrained_revision=json.loads(Path('/private/tmp/speech_sat_metadata.json').read_text())['sha']))
    vectorizer=DictVectorizer(dtype=np.float32)
    x=small_csr(vectorizer.fit_transform(feature_rows(train,sat_train,valid_only=True)))
    labels=np.array([v for r in train for v in r['labels'][:-1] if v>=0])
    if x.shape[0]!=len(labels):raise ValueError('Training feature/label misalignment')
    print(f'Lexical training: {x.shape[0]} gaps, {x.shape[1]} training-only features',flush=True)
    for alpha in [1e-5,1e-4]:
        name=f'lexical_joint_{alpha:g}'
        classifier=SGDClassifier(loss='log_loss',alpha=alpha,max_iter=40,tol=1e-4,average=True,random_state=42)
        classifier.fit(x,labels)
        p=lexical_probabilities(validation,sat_validation,vectorizer,classifier);choice=choose_threshold(validation,p)
        folder=WORK/name;folder.mkdir(exist_ok=True)
        joblib.dump(dict(vectorizer=vectorizer,classifier=classifier),folder/'model.joblib')
        record_experiment(name,dict(**metadata(name,protocol),status='complete',selected=choice,
            validation=metrics(validation,p,choice['threshold']),alpha=alpha,features=x.shape[1],
            training_iterations=int(classifier.n_iter_),artifact=str(folder)))
        print(name,'validation F1',choice['f1'],flush=True)


def bert(mask=False):
    splits,protocol=load_partitions();name='bert_joint_masked' if mask else 'bert_joint'
    pretrained_path=Path('/private/tmp/speech_bert_model_path.txt').read_text().strip()
    meta=dict(**metadata(name,protocol),pretrained_model='google/bert_uncased_L-4_H-256_A-4',
        pretrained_revision=json.loads(Path('/private/tmp/speech_bert_metadata.json').read_text())['sha'],
        augmentation='5% encoder-token masking; intact rule channel' if mask else 'none')
    start=time.monotonic()
    model,tokenizer,saved=fit_pretrained(splits['train'],splits['validation'],pretrained_path,WORK/name,meta,
        epochs=8,mask_probability=.05 if mask else 0.)
    p=pretrained_probabilities(model,tokenizer,splits['validation'])
    record_experiment(name,dict(**meta,status='complete',selected=saved['selected'],history=saved['history'],
        validation=metrics(splits['validation'],p,saved['threshold']),artifact=str(WORK/name),
        parameters=sum(v.numel() for v in model.parameters()),elapsed_seconds=time.monotonic()-start))


def select():
    """Persist choice before any new selected-model benchmark evaluation."""
    if (OUT/'improved_model.json').exists():raise ValueError('This round has already been benchmarked; do not reselect on its test outcomes')
    experiments=json.loads(EXPERIMENTS.read_text());complete={n:r for n,r in experiments.items() if r['status']=='complete'}
    if not complete:raise ValueError('No completed validation experiments')
    scores={n:r['validation']['overall']['f1'] for n,r in complete.items()}
    best=max(scores,key=scores.get)
    # The small/faster preference is only used inside the registered 0.001 band.
    eligible=[n for n,v in scores.items() if scores[best]-v<=.001]
    # The lexical model also needs SaT's much larger frozen backbone at runtime.
    # Count that dependency when applying the registered smaller/faster tie rule.
    order={'hybrid_overlap':0,'bert_joint':1,'bert_joint_masked':1,'sat_pretrained':2,
           'lexical_joint_0.0001':3,'lexical_joint_1e-05':3}
    best=min(eligible,key=lambda n:(order.get(n,4),-scores[n],n))
    result=dict(selected=best,criterion='Validation micro boundary F1; within 0.001 prefer smaller/faster',
        validation_scores=scores,selection_recorded_before_new_test_evaluation=True,
        input_sha256=complete[best]['input_sha256'],test_status=complete[best]['test_status'])
    save(EVIDENCE/'improvement_selection.json',result)
    print(json.dumps(result,indent=2));return best


def selected_probabilities(name,records):
    if name.startswith('bert'):
        model,tokenizer,_=load_bundle(WORK/name)
        return pretrained_probabilities(model,tokenizer,records)
    if name=='hybrid_overlap':return overlap_probabilities(records)
    if name=='sat_pretrained':return cache_probabilities('sat',records,sat_probabilities)
    if name.startswith('lexical'):
        import joblib
        saved=joblib.load(WORK/name/'model.joblib')
        sat=cache_probabilities('sat',records,sat_probabilities)
        return lexical_probabilities(records,sat,saved['vectorizer'],saved['classifier'])
    raise ValueError(f'Unknown selected model {name}')


def evaluate_selected():
    selection=json.loads((EVIDENCE/'improvement_selection.json').read_text());name=selection['selected']
    if (OUT/'improved_model.json').exists():
        saved=json.loads((OUT/'improved_model.json').read_text())
        if saved['selected']!=name:raise ValueError('Refusing to replace an evaluated selection')
        print('Returning the completed benchmark without reselecting or re-evaluating it')
        return saved
    experiments=json.loads(EXPERIMENTS.read_text());entry=experiments[name]
    splits,protocol=load_partitions(include_test=True)
    if selection['input_sha256']!=protocol['input_sha256']:raise ValueError('Selection/data hash mismatch')
    threshold=entry['selected']['threshold']
    test=splits['test'];p=selected_probabilities(name,test)
    result=dict(selected=name,threshold=threshold,validation=entry['validation'],test_status=protocol['test_status'],
        input_sha256=protocol['input_sha256'],test=metrics(test,p,threshold))
    target=[json.loads(s) for s in (EVIDENCE/'boundary_dataset.jsonl').read_text().splitlines()]
    q=selected_probabilities(name,target)
    result['original_asr']=metrics(target,q,threshold)
    result['original_asr_sha256']=digest(EVIDENCE/'boundary_dataset.jsonl')
    save(OUT/'improved_model.json',result)
    if name.startswith('bert'):
        destination=OUT/'segmenter';destination.mkdir(exist_ok=True)
        for filename in ['model.safetensors','tokenizer.json','segmenter.json']:
            shutil.copy2(WORK/name/filename,destination/filename)
    print('Selected-model reused benchmark F1:',result['test']['overall']['f1'],flush=True)
    print('Selected-model original ASR F1:',result['original_asr']['overall']['f1'],flush=True)
    return result


def bootstrap_comparison(records,new_rows,old_rows,seed=42):
    from speech_pipeline.boundary_model import total_counts
    people=sorted({r['participant'] for r in records})
    def array(rows,field):
        return np.array([[sum(r[field][k] for r in rows if r['participant']==p) for k in ('tp','fp','fn')] for p in people])
    a=array(new_rows,'counts');b=array(old_rows,'transformer')
    draws=np.random.default_rng(seed).integers(0,len(people),(10000,len(people)))
    def values(c):
        c=c[draws].sum(axis=1);den=2*c[:,0]+c[:,1]+c[:,2]
        return np.divide(2*c[:,0],den,out=np.zeros(len(c)),where=den>0)
    newer=values(a);older=values(b)
    return dict(method='10,000 paired participant-cluster resamples; descriptive on reused benchmark',
        participants=len(people),seed=seed,boundary_f1_ci=list(map(float,np.quantile(newer,[.025,.975]))),
        difference_vs_previous_joint_ci=list(map(float,np.quantile(newer-older,[.025,.975]))))


def downstream():
    """Final metric effects and official calculator checks; no model selection."""
    from tempfile import TemporaryDirectory
    import csv,spacy
    from openpyxl import load_workbook
    from speech_pipeline import utterances as auto
    from speech_pipeline.metrics import clinical_from_predictions
    from speech_pipeline.pipeline import lexical,clan_software_check
    result=json.loads((OUT/'improved_model.json').read_text());splits,protocol=load_partitions(include_test=True)
    records=splits['test']
    target=[json.loads(s) for s in (EVIDENCE/'boundary_dataset.jsonl').read_text().splitlines()]
    index={(r['corpus'],r['participant'],r['recording']):r for r in records+target}
    old=json.loads((OUT/'hybrid_model.json').read_text())
    if old['input_sha256']!=protocol['input_sha256']:raise ValueError('Old/new test source mismatch')
    result['uncertainty']=bootstrap_comparison(records,result['test']['recordings'],old['held_out_test']['recordings'])
    old_predictions={(r['corpus'],r['participant'],r['recording']):r['predicted_boundaries'] for group in
        [old['held_out_test'],old['additional_evaluation']] for r in group['recordings']}
    source=ROOT/'dataset/annotations/all_nafee_utterances.xlsx'
    provenance=json.loads((EVIDENCE/'boundary_dataset_sources.json').read_text())
    if provenance['sources'].get(str(source.relative_to(ROOT)))!=digest(source):raise ValueError('Human target annotations changed')
    workbook=load_workbook(source,read_only=True,data_only=True)
    human={(sheet.title,filename):[dict(text=auto.plain(t)) for t in texts] for sheet in workbook
           for filename,texts in auto.flatten_labels(sheet).items()};workbook.close()
    nlp=spacy.load('en_core_web_sm');table=[];coverage=[];checks=[];check_children={c:set() for c in ('ocsc','nyu_emerson')}
    def segmented(r,boundaries):
        words=r['surface_words'];cuts=[0,*boundaries,len(words)]
        return [dict(text=' '.join(words[a:b])) for a,b in zip(cuts,cuts[1:]) if a<b]
    with TemporaryDirectory(prefix='segmentation_metric_audit_') as temp:
        def scores(units):
            values={f'surface/{k}':v for k,v in clinical_from_predictions(units)['values'].items()}
            mlu=auto.word_mlu(units)
            values.update({f'utterances/{k}':mlu[k] for k in ('total_utterances','mlu_nutterances','mlu_words_sd')})
            if any(auto.analysis_text(u['text']) for u in units):
                try:
                    comparison=lexical(units,nlp,Path(temp)/'lexical.txt')
                    values.update({f'custom_normalized/{k}':v['custom'] for k,v in comparison.items()})
                    values.update({f'taaled/{k}':v['reference'] for k,v in comparison.items()})
                except (ValueError,ZeroDivisionError) as error:
                    coverage.append(dict(stage='lexical',status='unavailable',detail=str(error)))
            return values
        predictions=result['test']['recordings']+result['original_asr']['recordings']
        for at,prediction in enumerate(predictions):
            key=(prediction['corpus'],prediction['participant'],prediction['recording']);r=index[key]
            gold=segmented(r,r['gold_boundaries']) if r['corpus'] in ('ocsc','nyu_emerson') else human[r['participant'],r['recording']]
            truth=scores(gold)
            methods=dict(rules=segmented(r,[i+1 for i,v in enumerate(r['baseline'][:-1]) if v]),
                         previous_joint=segmented(r,old_predictions[key]),improved=segmented(r,prediction['predicted_boundaries']))
            for method,units in methods.items():
                mine=scores(units)
                for metric,reference in truth.items():
                    value=mine.get(metric)
                    table.append(dict(method=method,corpus=r['corpus'],participant=r['participant'],recording=r['recording'],
                        gem=r.get('gem'),metric=metric,reference=reference,predicted=value,
                        absolute_error=abs(value-reference) if value is not None and reference is not None else None,
                        absolute_percent_error=100*abs(value-reference)/abs(reference) if reference and value is not None else None,
                        status='available' if value is not None and reference is not None else 'unavailable_required_annotation_or_sample_size'))
            corpus=r['corpus']
            if corpus in check_children and len(check_children[corpus])<3 and r['participant'] not in check_children[corpus]:
                check_children[corpus].add(r['participant'])
                reference=clan_software_check(methods['improved'],Path(temp))
                checks.append(dict(corpus=corpus,participant=r['participant'],recording=r['recording'],clan=reference))
            if at%20==0:print(f'Final downstream metric comparison {at+1}/{len(predictions)}',flush=True)
    path=OUT/'improved_metric_errors.csv'
    with path.open('w',newline='',encoding='utf-8-sig') as f:
        writer=csv.DictWriter(f,fieldnames=list(table[0]));writer.writeheader();writer.writerows(table)
    save(EVIDENCE/'improved_calculator_checks.json',dict(scope='First three distinct held-out children per corpus; same predicted normalized input',checks=checks,coverage=coverage))
    import statistics
    summary={}
    for method in methods:
        summary[method]={}
        for metric in ('surface/MLUw','surface/TNW','surface/NDW','surface/TTR','surface/MATTR',
                       'utterances/total_utterances','taaled/basic_ntokens','taaled/basic_ntypes','taaled/simple_ttr_aw','taaled/mattr50_aw'):
            selected=[r for r in table if r['method']==method and r['metric']==metric and r['corpus'] in ('ocsc','nyu_emerson')]
            child={}
            for row in selected:
                if row['absolute_percent_error'] is not None:child.setdefault(row['participant'],[]).append(row['absolute_percent_error'])
            summary[method][metric]=dict(paragraphs=len(selected),available=sum(r['absolute_percent_error'] is not None for r in selected),
                median_paragraph_ape=statistics.median(r['absolute_percent_error'] for r in selected if r['absolute_percent_error'] is not None) if child else None,
                median_child_mean_ape=statistics.median(statistics.mean(v) for v in child.values()) if child else None)
    result['downstream']=summary
    result['metric_policy']='Main-body surface/filler/exclusion policy and normalized custom/TAALED policies reported separately; missing morphology/timing/grammar/maze labels remain unavailable'
    save(OUT/'improved_model.json',result)
    print('Final metric comparisons and participant uncertainty saved',flush=True)
    return result


def runtime_context():
    """Audit explicit pretrained workbook transfer after selection is frozen."""
    from openpyxl import load_workbook
    from speech_pipeline import utterances as auto
    from speech_pipeline.pipeline import difference
    result=json.loads((OUT/'improved_model.json').read_text())
    base=ROOT/'outputs/results/pipeline/uncorrected'
    report=json.loads((base/'comparison.json').read_text())
    if report['segmenter']['selected']!=result['selected']:raise ValueError('Runtime model differs from selection')
    index={(r['patient_id'],r['source_file']):r for r in report['recordings']}
    old_base=ROOT/'outputs/validation/method_comparison/analysis/uncorrected'
    old=json.loads((old_base/'comparison.json').read_text())
    old_index={(r['patient_id'],r['source_file']):r for r in old['recordings']}
    source=ROOT/'dataset/annotations/all_nafee_utterances.xlsx'
    expected=json.loads((EVIDENCE/'boundary_dataset_sources.json').read_text())['sources']
    if expected[str(source.relative_to(ROOT))]!=digest(source):raise ValueError('Human labels changed')
    workbook=load_workbook(source,read_only=True,data_only=True);rows=[]
    for sheet in workbook:
        for filename,texts in auto.flatten_labels(sheet).items():
            key=sheet.title,filename;r=index[key];previous=old_index[key]
            raw=ROOT/'dataset/uncorrected'/sheet.title/filename
            if not digest(raw)==r['raw_sha256']==previous['raw_sha256']:raise ValueError('Runtime input changed')
            gold=[dict(text=auto.plain(t)) for t in texts]
            units=json.loads((base/r['utterance_file']).read_text())['utterances']
            earlier=json.loads((old_base/previous['utterance_file']).read_text())['utterances']
            if auto.stream(units)[0]!=auto.stream(earlier)[0]:raise ValueError('The segmenter altered source words')
            counts=auto.boundary_counts(gold,units);baseline=auto.boundary_counts(gold,earlier)
            reference=auto.word_mlu(gold);mine=auto.word_mlu(units)
            rows.append(dict(participant=sheet.title,recording=filename,counts=counts,previous_counts=baseline,
                scores=auto.boundary_scores(counts),mlu={k:difference(mine[k],v) for k,v in reference.items()
                    if isinstance(v,(int,float)) or v is None}))
    workbook.close()
    def aggregate(field):
        return auto.boundary_scores({k:sum(r[field][k] for r in rows) for k in rows[0][field]})
    result['runtime_context']=dict(scope='Descriptive original ASR cohort; both methods use observed speaker turns, no individual source-row cuts',
        recordings=rows,overall=aggregate('counts'),previous=aggregate('previous_counts'))
    checks=[]
    for kind in ('uncorrected','corrected'):
        path=ROOT/f'outputs/results/pipeline/{kind}/comparison.json';run=json.loads(path.read_text())
        if run['segmenter']['selected']!=result['selected']:raise ValueError('Runtime selection mismatch')
        for record in run['recordings']:
            comparisons=record['clan']['requested_metrics']
            for metric,value in comparisons.items():
                error=value['absolute_error'];tolerance=0 if metric in ('TNW','NDW') else .000501
                if error is not None and error>tolerance:raise ValueError(f'CLAN mismatch {kind}/{record["source_file"]}: {metric} {error}')
            checks.append(dict(dataset=kind,participant=record['patient_id'],recording=record['source_file'],
                max_clan_error=max((v['absolute_error'] for v in comparisons.values() if v['absolute_error'] is not None),default=0),
                taaled_measures=len(record['lexical'])))
    result['runtime_metric_checks']=dict(records=len(checks),all_clan_within_display_precision=True,
        taaled_measures_per_recording=sorted({r['taaled_measures'] for r in checks}),checks=checks)
    save(OUT/'improved_model.json',result)
    print('Full pipeline turn-aware F1:',result['runtime_context']['overall']['boundary_f1'],flush=True)
    print('Official CLAN checks passed for',len(checks),'recordings',flush=True)
    return result


def refresh_calculators():
    """Refresh exported words/calculators without changing saved boundaries."""
    from tempfile import TemporaryDirectory
    from speech_pipeline import utterances as auto
    from speech_pipeline.pipeline import clan_software_check
    from speech_pipeline.export import export_all
    with TemporaryDirectory(prefix='clan_export_refresh_') as temp:
        for kind in ('uncorrected','corrected'):
            base=ROOT/f'outputs/results/pipeline/{kind}';path=base/'comparison.json'
            report=json.loads(path.read_text())
            for record in report['recordings']:
                units=json.loads((base/record['utterance_file']).read_text())['utterances']
                previous=base/record['chat_file']
                if digest(previous)!=record['chat_sha256']:raise ValueError('Saved CHAT input changed')
                body=auto.chat_text(units,f'{record["patient_id"]}/{record["source_file"]}; speaker {record["target_speaker"]}')
                sha=hashlib.sha256(body.encode()).hexdigest()
                stem=previous.stem.rsplit('_auto_',1)[0]
                current=previous.with_name(f'{stem}_auto_{sha[:12]}.cha');current.write_text(body)
                record.update(chat_file=str(current.relative_to(base)),chat_sha256=sha,
                    clan=clan_software_check(units,Path(temp)))
                if current!=previous:previous.unlink()
            report['calculator_refresh']=dict(reason='Canonical apostrophes in CHAT match existing clinical surface policy; saved predictions unchanged',
                code_sha256={p:digest(ROOT/p) for p in ['speech_pipeline/utterances.py','speech_pipeline/pipeline.py']})
            save(path,report);export_all(report,base)
            print('Refreshed actual CLAN checks:',kind,len(report['recordings']),flush=True)
        checks=json.loads((EVIDENCE/'improved_calculator_checks.json').read_text())
        result=json.loads((OUT/'improved_model.json').read_text())
        splits,_=load_partitions(include_test=True)
        index={(r['corpus'],r['participant'],r['recording']):r for r in splits['test']}
        predicted={(r['corpus'],r['participant'],r['recording']):r for r in result['test']['recordings']}
        for check in checks['checks']:
            key=check['corpus'],check['participant'],check['recording'];r=index[key]
            cuts=[0,*predicted[key]['predicted_boundaries'],len(r['surface_words'])]
            units=[dict(text=' '.join(r['surface_words'][a:b])) for a,b in zip(cuts,cuts[1:]) if a<b]
            check['clan']=clan_software_check(units,Path(temp))
        save(EVIDENCE/'improved_calculator_checks.json',checks)


def improvement_report():
    """Latest results section and one overview figure; no model evaluation."""
    path=OUT/'improved_model.json'
    if not path.exists():return []
    result=json.loads(path.read_text());experiments=json.loads(EXPERIMENTS.read_text())
    complete={k:v for k,v in experiments.items() if v['status']=='complete'}
    old=json.loads((OUT/'hybrid_model.json').read_text());test=result['test']['overall']
    lines=['## Earlier round: four-layer pretrained context plus rules','',
        'The best of the registered paragraph candidates is a four-layer pretrained BERT encoder with a learned rule embedding and a left/right word-gap head. It uses overlapping 256-subtoken windows and center-weighted logit aggregation. The selected checkpoint is epoch 7, threshold 0.39; 11,238,562 parameters. One set of word-preserving segments feeds all metrics. Plain-paragraph analysis loads this local model automatically; diarized workbook processing retains the established rules.',
        '', 'Both downloaded corpora contribute training data: 301 usable OCSC children and 30 NYU–Emerson children; 907 paragraphs and 347,502 words. The frozen partitions contain 233 training, 49 validation and 49 test children (641/140/126 paragraphs). Published punctuation, casing, row breaks, speaker turns, morphology and timestamps are withheld from paragraph inputs. No child crosses partitions. Weight fitting, candidate selection, epochs and thresholds use training/validation only.',
        '', '**Evaluation status:** these are reused research benchmarks, previously examined in earlier work. The candidate plan was recorded before these experiments and the selection was saved before the selected model was evaluated. This is not a fresh confirmatory test. The subsequent sections preserve the earlier comparisons; references there to initially untouched test children describe that earlier stage.',
        '', 'The [literature review](../../../speech_pipeline/research/SEGMENTATION_RESEARCH.md) covers 14 primary papers. Pretrained bidirectional context, domain adaptation and overlapping inference motivated the implemented candidates. Only locally measured scores are reported. Recent LLM and acoustic approaches discussed in the review were not run.',
        '', '| Registered candidate | Validation boundary F1 |','| --- | ---: |']
    for name,entry in complete.items():lines.append(f'| {name} | {100*entry["validation"]["overall"]["f1"]:.2f}% |')
    lines += ['', 'The 5% token-masking candidate did not improve validation F1. An initial sparse-model run failed on CSR index width; it was repaired and both unchanged candidate configurations completed. These bounded single-seed comparisons do not prove a globally optimal architecture or an isolated causal benefit of each component. Candidate checkpoints and caches stay in temporary storage; only the selected production bundle is retained in the project.',
        '', '| Evaluation | Previous joint model F1 | Selected model F1 |','| --- | ---: | ---: |']
    for corpus,value in result['test']['by_corpus'].items():
        lines.append(f'| {corpus} ({value["participants"]} children) | {100*old["held_out_test"]["by_corpus"][corpus]["transformer"]["f1"]:.2f}% | {100*value["f1"]:.2f}% |')
    for label,new,key in [('All TalkBank test paragraphs',test,'held_out_test'),('Original ASR continuous paragraphs',result['original_asr']['overall'],'additional_evaluation')]:
        lines.append(f'| {label} | {100*old[key]["overall"]["transformer"]["f1"]:.2f}% | {100*new["f1"]:.2f}% |')
    exact=test['exact']
    lines += ['',f'TalkBank test precision **{100*test["precision"]:.2f}%**, recall **{100*test["recall"]:.2f}%**. Boundary errors: {test["counts"]["fp"]:,} extra and {test["counts"]["fn"]:,} missed. Strict complete-utterance span F1 is **{100*exact["f1"]:.2f}%** ({exact["matched"]:,} exactly matched spans of {exact["reference"]:,} reference spans); this stricter score includes gaps masked from boundary evaluation around unintelligible words.',
        '', '![Validation candidates and corpus results](segmentation_improvement.png)']
    lines += ['', '| Test task | Children | Selected boundary F1 |','| --- | ---: | ---: |']
    for task,value in result['test']['by_task'].items():
        lines.append(f'| {task} | {value["participants"]} | {100*value["f1"]:.2f}% |')
    lines += ['', 'Five lowest-scoring test children (all child and recording results remain in the JSON):','',
        '| Child | Paragraphs | Extra boundaries | Missed boundaries | Boundary F1 |','| --- | ---: | ---: | ---: | ---: |']
    for child,value in sorted(result['test']['by_child'].items(),key=lambda item:item[1]['f1'])[:5]:
        lines.append(f'| {child} | {value["records"]} | {value["counts"]["fp"]} | {value["counts"]["fn"]} | {100*value["f1"]:.2f}% |')
    if 'uncertainty' in result:
        u=result['uncertainty'];lines += ['',f'Paired 10,000-resample child-cluster bootstrap: test F1 interval {100*u["boundary_f1_ci"][0]:.2f}–{100*u["boundary_f1_ci"][1]:.2f}%; improvement versus the previous joint model {100*u["difference_vs_previous_joint_ci"][0]:.2f}–{100*u["difference_vs_previous_joint_ci"][1]:.2f} percentage points. These intervals are descriptive on this reused benchmark.']
    if 'downstream' in result:
        lines += ['', '| Surface measure | Rules median child mean APE | Previous joint | Selected |','| --- | ---: | ---: | ---: |']
        for metric in ('MLUw','TNW','NDW','TTR','MATTR'):
            values=[result['downstream'][m]['surface/'+metric]['median_child_mean_ape'] for m in ('rules','previous_joint','improved')]
            lines.append('| '+metric+' | '+' | '.join(f'{v:.2f}%' if v is not None else 'unavailable' for v in values)+' |')
        lines += ['', '`improved_metric_errors.csv` contains every recording/paragraph comparison, all 17 TAALED measures and explicit unavailable statuses. Surface counts and normalized TAALED measures have different lexical policies. Segmentation alone does not improve TNW/NDW when words and eligibility remain identical; changed exclusion of unintelligible material can affect them. Six predicted public paragraphs from distinct test children also have actual same-input CLAN checks in `improved_calculator_checks.json`.']
    if 'runtime_context' in result:
        r=result['runtime_context'];lines += ['', f'**Transfer failure:** with observed speaker turns, the pretrained model scored {100*r["overall"]["boundary_f1"]:.2f}% versus {100*r["previous"]["boundary_f1"]:.2f}% for the established rules. This context differs from continuous paragraphs. Word-gap alignment coverage is {100*r["overall"]["reference_gap_coverage"]:.2f}% of reference gaps. Predictions preserve the original word stream. The workbook default therefore remains the established rules; explicit `--segmenter improved` is experimental. This deployment safeguard was informed by the descriptive cohort result and is not a newly validation-selected ensemble. No weights or thresholds were adjusted on these cohort labels.']
    if 'runtime_metric_checks' in result:
        lines += ['', 'The pretrained transfer trial covered all 37 recordings in both transcript conditions. All 74 actual CLAN MLU/FREQ comparisons match counts exactly and ratios within three-decimal display precision; 17 custom/TAALED comparisons were computed per recording. Exact trial CHAT inputs and raw official outputs remain in `improved_calculator_checks.json`. One corrected transcript exposed a curly/straight-apostrophe type-count mismatch; CHAT export now matches the existing surface normalization, without changing source text or boundaries. Canonical `outputs/results/pipeline/` outputs use the retained workbook default. Final paper, poster and abstract are unchanged.']
    if 'production_verification' in result:
        v=result['production_verification'];lines += ['', f'The retained workbook default was rerun and verified separately: {v["checked_recordings"]} recording-condition results, all actual CLAN checks within output precision, and {v["taaled_measures_per_recording"]} TAALED comparisons each. Its original-cohort F1 remains {100*v["actual_original_cohort_scores"]["boundary_f1"]:.2f}%. The final regression suite passed 44 tests. Source hashes for all 907 public paragraphs and the final publication-file hashes were checked unchanged.']
    if 'publication_comparison' in result:
        import csv
        comparison=result['publication_comparison']
        def published(filename,group,column):
            with (ROOT/'outputs/results/paper'/filename).open(encoding='utf-8-sig') as f:
                return {r['metric']:float(r[column]) for r in csv.DictReader(f) if r['comparison']==group}
        paper=published('publication_tables_median.csv','Nafee complete participant sample','median_participant_error_percent')
        poster=published('publication_tables.csv','Nafee recording segments','mape')
        lines += ['', '## Direct comparison with the paper and poster','',
            'The frozen pretrained model did not improve the original 10-participant full pipeline. The table uses the same primary human labels and speaker-turn context, retaining each publication\'s aggregation. Lower errors are better.', '',
            '| Measure | Paper: previous → pretrained (complete-sample median APE) | Poster: previous → pretrained (participant-balanced mean recording APE) |',
            '| --- | ---: | ---: |']
        for label,key in [('Utterance count','total_utterances'),('MLUw','mlu_words'),('Length SD','mlu_words_sd')]:
            metric={'mlu_words':'MLUw','mlu_words_sd':'SD'}.get(key,key)
            lines.append(f'| {label} | {paper[key]:.2f}% → {comparison["pretrained_paper_mdape"][metric]:.2f}% | {poster[key]:.2f}% → {comparison["pretrained_poster_mape"][metric]:.2f}% |')
    if 'final_cohort_policy' in result:
        lines += ['', 'The original **10-participant cohort is final evaluation only going forward**: no training, adaptation, threshold/epoch selection, rule refinement, or model/ensemble/routing selection. Every choice must be frozen using TalkBank training/validation before testing. The reported cutoff was selected on TalkBank validation, rather than manually guessed, and remains frozen. Earlier rule development and a deployment-routing decision already used this cohort; it cannot retrospectively become an untouched independent final test. No weights, thresholds or predictions were changed for the publication comparison.']
    if 'reassessment' in result:
        lines[2:2]=['**Historical section:** all selected-model scores, cutoff values, runtime checks and publication comparisons in this section describe the earlier four-layer BERT round. The reassessment above supersedes its plain-paragraph model. The current canonical pipeline outputs come from the explicitly requested paragraph mode; the application default for speaker-turn workbooks remains rules. Statements below about the then-current canonical outputs and 44-test suite refer to that earlier run.', '']
    lines += ['', '**Metric requirements:** MLUw, TNW, NDW, TTR, MATTR-50 and the existing TAALED measures remain available. The reviewed-annotation API supports MLUm, WPM, PGU, omission/substitution counts and maze percentage when complete supporting morphology, timing or labels are supplied. The paragraph model does not infer those annotations. Missing evidence stays missing. CHAT utterances are not automatically SALT C-units, and no SALT software equivalence or clinical interchangeability is claimed.',
        '', 'Per-child/task results and every predicted cut are in `improved_model.json`; training/selection evidence is under `outputs/validation/benchmark/`. A new independent child-speech cohort is needed for confirmatory evaluation.', '']
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    fig,axes=plt.subplots(1,2,figsize=(12,5),dpi=180)
    names=list(complete);scores=[100*complete[n]['validation']['overall']['f1'] for n in names]
    axes[0].barh(names,scores,color=['#237a59' if n==result['selected'] else '#417ca0' for n in names]);axes[0].set(xlabel='Validation boundary F1 (%)',xlim=(0,100))
    corpora=list(result['test']['by_corpus']);x=np.arange(len(corpora))
    axes[1].bar(x-.18,[100*old['held_out_test']['by_corpus'][c]['transformer']['f1'] for c in corpora],.36,label='Previous joint')
    axes[1].bar(x+.18,[100*result['test']['by_corpus'][c]['f1'] for c in corpora],.36,label='Selected pretrained joint',color='#237a59')
    axes[1].set(xticks=x,xticklabels=corpora,ylabel='Reused test boundary F1 (%)',ylim=(0,100));axes[1].legend()
    fig.tight_layout();fig.savefig(OUT/'segmentation_improvement.png');plt.close(fig)
    return lines


def report():
    from speech_pipeline.research.validation import write_report
    write_report(json.loads((OUT/'existing_cohort.json').read_text()))


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('stage',choices=['fetch-models','baselines','bert','bert-masked','select','evaluate','downstream','refresh-calculators','runtime-context','report'])
    args=parser.parse_args()
    {'fetch-models':fetch_models,'baselines':baselines,'bert':bert,'bert-masked':lambda:bert(True),'select':select,
     'evaluate':evaluate_selected,'downstream':downstream,'refresh-calculators':refresh_calculators,
     'runtime-context':runtime_context,'report':report}[args.stage]()
